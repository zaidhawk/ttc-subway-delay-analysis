"""
Builds an Excel dashboard from the cleaned data.
Run clean_data.py first. The file is saved to output/.
"""
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.axis import ChartLines
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

CLEAN = Path("data/clean")
OUT = Path("output/TTC_Subway_Delay_Dashboard.xlsx")

MAIN_LINES = ["Line 1 Yonge-University", "Line 2 Bloor-Danforth", "Line 4 Sheppard"]
COLORS = ["2A78D6", "EB6834", "1BAF7A"]  # one colour per line
HEADER_FILL = PatternFill("solid", fgColor="1F3864")


def light_gridlines():
    return ChartLines(spPr=GraphicalProperties(ln=LineProperties(solidFill="E3E3E3")))


def summarize(df, by):
    # Count incidents, delays and delay minutes for each group
    g = df.groupby(by)
    out = pd.DataFrame({
        "Incidents": g.size(),
        "Incidents Causing Delay": g["Caused Delay"].sum(),
        "Total Delay (min)": g["Min Delay"].sum(),
    })
    out["Avg Delay When Delayed (min)"] = (
        out["Total Delay (min)"] / out["Incidents Causing Delay"].where(lambda s: s > 0)
    ).round(1)
    out["% of All Delay Minutes"] = (out["Total Delay (min)"] / df["Min Delay"].sum()).round(3)
    return out.sort_values("Total Delay (min)", ascending=False).reset_index()


def build_tables(df):
    months = df["Month"].dt.strftime("%Y-%m")

    by_month = (
        df[df["Line"].isin(MAIN_LINES)]
        .pivot_table(index=months, columns="Line", values="Min Delay", aggfunc="sum", fill_value=0)
        .reindex(columns=MAIN_LINES)
    )
    by_month.index.name = "Month"
    by_month["All Lines Incidents"] = df.groupby(months).size()
    by_month = by_month.reset_index()

    by_cause = summarize(df, ["Code", "Code Description", "Category"]).head(20)

    by_station = summarize(df, "Station").head(15)

    # Train equipment failures
    eu = df[df["Category"] == "Equipment (Rolling Stock)"]
    rolling_stock = summarize(eu, ["Code", "Code Description"])
    rolling_stock_by_line = summarize(eu[eu["Line"].isin(MAIN_LINES)], "Line")
    rolling_stock_by_line["Incidents per Month"] = (
        rolling_stock_by_line["Incidents"] / df["Month"].nunique()
    ).round(1)

    by_hour = summarize(df, "Hour").sort_values("Hour")

    return {
        "By Line": summarize(df, "Line"),
        "By Category": summarize(df, "Category"),
        "By Month": by_month,
        "Top Causes": by_cause,
        "Top Stations": by_station,
        "Rolling Stock": rolling_stock,
        "Rolling Stock by Line": rolling_stock_by_line,
        "By Hour": by_hour,
    }


def kpis(df):
    delayed = df[df["Caused Delay"]]
    return [
        ("Period", f"{df['Date'].min():%b %Y} - {df['Date'].max():%b %Y}"),
        ("Total incidents logged", len(df)),
        ("Incidents that caused a delay", len(delayed)),
        ("% of incidents causing delay", round(len(delayed) / len(df), 3)),
        ("Total delay (hours)", round(df["Min Delay"].sum() / 60)),
        ("Average delay when delayed (min)", round(delayed["Min Delay"].mean(), 1)),
        ("Median delay when delayed (min)", delayed["Min Delay"].median()),
        ("Delays of 60+ minutes", int(df["Long Delay Flag"].sum())),
        ("Rolling stock (EU) incidents", int((df["Category"] == "Equipment (Rolling Stock)").sum())),
    ]


def style_sheet(ws, header_row=1, freeze=True):
    for cell in ws[header_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for col in ws.columns:
        width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(width + 2, 10), 60)
    for row in ws.iter_rows(min_row=header_row + 1):
        for cell in row:
            header = ws.cell(header_row, cell.column).value or ""
            if header.startswith("%"):
                cell.number_format = "0.0%"
            elif isinstance(cell.value, int):
                cell.number_format = "#,##0"
            elif isinstance(cell.value, float):
                cell.number_format = "#,##0.0"
    if freeze:
        ws.freeze_panes = ws.cell(header_row + 1, 1)


def bar_chart(ws, title, cats_col, vals_col, n_rows, horizontal=True, color=COLORS[0]):
    chart = BarChart()
    chart.type = "bar" if horizontal else "col"
    chart.title = title
    chart.title.overlay = False
    chart.legend = None
    data = Reference(ws, min_col=vals_col, min_row=1, max_row=n_rows + 1)
    cats = Reference(ws, min_col=cats_col, min_row=2, max_row=n_rows + 1)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.series[0].graphicalProperties.solidFill = color
    chart.series[0].graphicalProperties.line.noFill = True
    chart.gapWidth = 60
    if horizontal:
        chart.x_axis.scaling.orientation = "maxMin"  # biggest bar on top
        chart.y_axis.crosses = "max"  # keep value axis at the bottom
    chart.y_axis.majorGridlines = light_gridlines()
    chart.y_axis.numFmt = "#,##0"
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    chart.width, chart.height = 16, 8.5
    return chart


def monthly_chart(ws, n_rows):
    chart = LineChart()
    chart.title = "Monthly Delay Minutes by Line"
    chart.title.overlay = False
    chart.y_axis.majorGridlines = light_gridlines()
    chart.y_axis.numFmt = "#,##0"
    data = Reference(ws, min_col=2, max_col=1 + len(MAIN_LINES), min_row=1, max_row=n_rows + 1)
    cats = Reference(ws, min_col=1, min_row=2, max_row=n_rows + 1)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    for series, color in zip(chart.series, COLORS):
        series.graphicalProperties.line.solidFill = color
        series.graphicalProperties.line.width = 22000
        series.smooth = False
    chart.legend.position = "b"
    chart.legend.overlay = False
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    chart.x_axis.tickLblSkip = 3
    chart.width, chart.height = 33, 9
    return chart


def main():
    df = pd.read_csv(CLEAN / "subway_delays_clean.csv", parse_dates=["Date", "Month"])
    quality = pd.read_csv(CLEAN / "data_quality_log.csv")
    tables = build_tables(df)

    OUT.parent.mkdir(exist_ok=True)
    with pd.ExcelWriter(OUT, engine="openpyxl") as xw:
        pd.DataFrame(kpis(df), columns=["KPI", "Value"]).to_excel(
            xw, sheet_name="Dashboard", index=False, startrow=2)
        for name, table in tables.items():
            table.to_excel(xw, sheet_name=name, index=False)
        quality.to_excel(xw, sheet_name="Data Quality", index=False)
        df.assign(Date=df["Date"].dt.date, Month=df["Month"].dt.strftime("%Y-%m")).to_excel(
            xw, sheet_name="Clean Data", index=False)

    wb = load_workbook(OUT)

    for name in list(tables) + ["Data Quality", "Clean Data"]:
        style_sheet(wb[name])

    # Make the clean data an Excel table so it can be used in pivot tables
    cd = wb["Clean Data"]
    cd.add_table(Table(displayName="Delays", ref=cd.dimensions,
                       tableStyleInfo=TableStyleInfo(name="TableStyleLight9", showRowStripes=True)))

    # Dashboard sheet
    ws = wb["Dashboard"]
    ws["A1"] = "TTC Subway Delay Dashboard"
    ws["A1"].font = Font(bold=True, size=18, color="1F3864")
    ws["A2"] = "Source: City of Toronto Open Data - TTC Subway Delay Data. Built with Python (pandas + openpyxl)."
    ws["A2"].font = Font(italic=True, color="666666")
    style_sheet(ws, header_row=3, freeze=False)
    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 22
    for row in ws.iter_rows(min_row=4, max_col=2):
        row[1].alignment = Alignment(horizontal="right")
        if row[0].value and row[0].value.startswith("%"):
            row[1].number_format = "0.0%"

    n = lambda name: len(tables[name])
    ws.add_chart(monthly_chart(wb["By Month"], n("By Month")), "D3")
    ws.add_chart(bar_chart(wb["By Line"], "Delay Minutes by Line", 1, 4, len(MAIN_LINES)), "A23")
    ws.add_chart(bar_chart(wb["By Category"], "Delay Minutes by Cause Category", 1, 4, 5), "J23")
    ws.add_chart(bar_chart(wb["Top Causes"], "Top 10 Causes by Delay Minutes", 2, 6, 10), "A41")
    ws.add_chart(bar_chart(wb["Rolling Stock"], "Rolling Stock Failures - Top 10 by Delay Minutes",
                           2, 5, min(10, n("Rolling Stock"))), "J41")

    wb.move_sheet("Dashboard", offset=-wb.index(wb["Dashboard"]))
    wb.active = 0
    wb.save(OUT)
    print(f"Saved {OUT}")


if __name__ == "__main__":
    main()
